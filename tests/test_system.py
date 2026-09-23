import copy
import io
import json
import time
from concurrent.futures import ThreadPoolExecutor
from types import SimpleNamespace
from uuid import uuid4
import httpx
import pytest
from fastapi.testclient import TestClient
from backend.server import create_app
from backend.catalog import Catalog, CatalogError
from backend.cart import CartError
from backend.normalize import normalize
from backend.attachments import extract, extract_document
from agent.analogs import rank_analogs

@pytest.fixture
def system(tmp_path,monkeypatch):
    monkeypatch.setenv('AGENT_MODE','demo')
    monkeypatch.setenv('CATALOG_MODE','snapshot')
    monkeypatch.setenv('EKT_SELLABLE_STORE_IDS','')
    cat=Catalog()
    app=create_app(cat,tmp_path/'db.sqlite')
    client=TestClient(app)
    csrf=client.get('/api/session').json()['csrf']
    client.headers['X-CSRF-Token']=csrf
    return client,cat,app.state.cart,app

def chat(client,text,**kw):
    return client.post('/api/chat',json={'message':text,'request_id':str(uuid4()),**kw})

def propose(client,pid='demo-2',qty=2):
    response=client.post('/api/proposals',json={'product_id':pid,'quantity':qty})
    assert response.status_code==200,response.text
    return response.json()

def confirm(client,p): return client.post('/api/proposals/'+p['id']+'/confirm',json={'confirm':True})

def test_existing_article_and_conflict(system):
    c,cat,_,_=system
    r=chat(c,'Покажи 027228').json()
    p=r['products'][0]
    assert p['id']=='515291' and p['price']==64920
    assert p['specs']['current'] is None
    assert next(x for x in p['conflicts'] if x['field']=='current')['values']==[160,250]
    assert p['currency']=='KZT' and p['stock'] is None
    assert '160' in r['answer'] and '250' in r['answer']
    assert not rank_analogs(p,list(cat.products.values()))['items']

def test_spec_search_and_ambiguity(system):
    c,*_=system
    r=chat(c,'автомат 16 А').json()
    assert len(r['products'])>=2 and 'несколько' in r['answer']
    assert all(p['specs']['current']==16 for p in r['products'])

def test_zero_stock_analog(system):
    c,*_=system
    r=chat(c,'DEMO-A16').json()
    assert r['products'][0]['stock']==0
    assert [x['product']['id'] for x in r['analogs']]==['demo-2']
    assert r['analogs'][0]['matches'] and r['analogs'][0]['unknown']

def test_no_cart_before_confirmation_and_cancel(system):
    c,*_=system
    p=propose(c)
    assert c.get('/api/cart').json()['items']==[]
    c.post('/api/proposals/'+p['id']+'/cancel',json={})
    assert confirm(c,p).status_code==409
    assert c.get('/api/cart').json()['items']==[]

def test_success_replay_and_cart_page(system):
    c,*_=system
    p=propose(c)
    assert confirm(c,p).json()['cart_url']=='/cart'
    assert confirm(c,p).json()['already_confirmed']
    assert c.get('/api/cart').json()['items'][0]['quantity']==2
    assert c.get('/cart').status_code==200

def test_changed_stock_invalidates_proposal(system):
    c,cat,*_=system
    p=propose(c)
    cat.products['demo-2']['stock']=1
    assert confirm(c,p).status_code==409
    assert c.get('/api/cart').json()['items']==[]
    cat.products['demo-2']['stock']=5
    assert confirm(c,p).status_code==409

def test_changed_price_and_sufficient_stock_still_require_reconfirm(system):
    c,cat,*_=system
    p=propose(c);cat.products['demo-2']['price']=1800
    assert confirm(c,p).status_code==409
    p=propose(c);cat.products['demo-2']['stock']=4
    assert confirm(c,p).status_code==409

def test_concurrent_confirmation_exactly_once(system):
    c,_,cart,_=system
    p=propose(c)
    sid=c.cookies['ekt_session']
    with ThreadPoolExecutor(max_workers=8) as pool:
        results=list(pool.map(lambda _:cart.confirm(sid,p['id']),range(12)))
    assert sum(not r['already_confirmed'] for r in results)==1
    assert cart.get_cart(sid)['items'][0]['quantity']==2

def test_existing_cart_counts_toward_stock(system):
    c,*_=system
    confirm(c,propose(c,qty=4))
    assert c.post('/api/proposals',json={'product_id':'demo-2','quantity':2}).status_code==409

def test_user_isolation_and_csrf(system):
    c,_,_,app=system
    p=propose(c)
    b=TestClient(app); b.headers['X-CSRF-Token']=b.get('/api/session').json()['csrf']
    assert confirm(b,p).status_code==409
    assert b.get('/api/cart').json()['items']==[]
    confirm(c,p)
    assert b.get('/api/cart').json()['items']==[]
    assert b.post('/api/proposals',json={'product_id':'demo-2','quantity':1},headers={'X-CSRF-Token':'bad'}).status_code==403

def test_expiry_replacement_and_invalid_confirmation(system):
    c,_,cart,_=system
    old=propose(c);new=propose(c,qty=1)
    assert confirm(c,old).status_code==409
    with cart.connection() as db: db.execute('UPDATE proposals SET expires=0 WHERE id=?',(new['id'],))
    assert confirm(c,new).status_code==409
    p=propose(c)
    assert c.post('/api/proposals/'+p['id']+'/confirm',json={'confirm':'true'}).status_code==422
    assert c.post('/api/proposals/'+p['id']+'/confirm',json={'confirm':False}).status_code==400

def test_text_confirmation_bound_and_retry(system):
    c,*_=system
    r=chat(c,'Добавь 2 штуки DEMO-B16').json()
    assert r['proposal']
    assert not chat(c,'да добавь').json()['cart_url']
    body={'message':'Да, добавь','proposal_id':r['proposal']['id'],'request_id':str(uuid4())}
    first=c.post('/api/chat',json=body)
    assert first.status_code==200 and first.json()['cart_url']=='/cart'
    assert c.post('/api/chat',json=body).json()==first.json()
    assert c.get('/api/cart').json()['items'][0]['quantity']==2

def test_stores_and_unknowns():
    raw={'id':1,'name':'Test','quantity':25,'stores':[{'id':1,'name':'Брак','quantity':20},{'id':2,'name':'Алматы','quantity':5}]}
    assert normalize(raw,{'kind':'test'},[])['stock'] is None
    p=normalize(raw,{'kind':'test'},[2]);assert p['stock']==5 and p['quantity']==25
    assert p['price'] is None and p['specs']['current'] is None

def test_unit_conversion_and_ranges():
    from backend.normalize import number
    assert number('20…80 V') is None
    assert number('0.05-5 V') is None
    p=normalize({'id':1,'name':'Test','properties':{'NOMINALNOE_NAPRYAZHENIE':'0,4 кВ','NOMINALNYY_TOK':'500 мА','NOMINALNAYA_OTKLYUCHAYUSHCHAYA_SPOSOBNOST':'18000 А'}},{'kind':'test'})
    assert p['specs']['voltage']==400 and p['specs']['current']==.5 and p['specs']['breaking']==18

def test_api_failure_no_cart_change(system,monkeypatch):
    c,cat,*_=system
    p=propose(c)
    monkeypatch.setattr(cat,'get',lambda *a,**kw: (_ for _ in ()).throw(CatalogError('API каталога недоступен.')))
    assert confirm(c,p).status_code==503
    assert c.get('/api/cart').json()['items']==[]

def test_snapshot_cannot_be_purchased_and_terms_unknown(system):
    c,*_=system
    assert c.post('/api/proposals',json={'product_id':'515291','quantity':1}).status_code==409
    assert 'не предоставлены' in chat(c,'условия доставки').json()['answer']

def test_pagination_repeat_and_count(tmp_path):
    calls=[]
    def respond(request):
        calls.append(str(request.url))
        if 'detail' in request.url.path: return httpx.Response(200,json={'id':1,'name':'Test'})
        return httpx.Response(200,json={'page':1,'count':1,'per_page':1,'items':[{'id':1,'name':'Test'}]})
    client=httpx.Client(base_url='https://ekt.kz',transport=httpx.MockTransport(respond))
    cat=Catalog(tmp_path,client)
    coverage=cat.import_pages(5)
    assert coverage['real_products']==1 and coverage['stop_reason']=='повторяющаяся страница' and not coverage['complete']
    length=len(calls);cat.import_pages(5);assert len(calls)==length

def test_partial_detail_failure_and_list_failure(tmp_path):
    def respond(r):
        if 'detail' in r.url.path:return httpx.Response(503)
        return httpx.Response(200,json={'items':[{'id':7,'name':'List only'}] if r.url.params['page']=='1' else []})
    cat=Catalog(tmp_path,httpx.Client(base_url='https://ekt.kz',transport=httpx.MockTransport(respond)))
    assert cat.import_pages(3)['detail_errors']==1
    assert cat.get('7')['stock'] is None
    fail=Catalog(tmp_path/'new',httpx.Client(base_url='https://ekt.kz',transport=httpx.MockTransport(lambda r:httpx.Response(401))))
    with pytest.raises(CatalogError):fail.import_pages()

def test_unknown_critical_parameter_blocks_analog(system):
    _,cat,*_=system
    p=copy.deepcopy(cat.get('demo-1'));p['specs']['breaking']=None
    assert rank_analogs(p,list(cat.products.values()))['items']==[]

def test_excel_word_and_pdf_extraction():
    from openpyxl import Workbook
    from docx import Document
    from pypdf import PdfWriter
    wb=Workbook();wb.active.append(['Артикул','Количество']);wb.active.append(['DEMO-B16',2]);b=io.BytesIO();wb.save(b)
    assert extract_document('.xlsx',b.getvalue())[0]['quantity']==2
    doc=Document();doc.add_paragraph('DEMO-B16 | 3');b=io.BytesIO();doc.save(b)
    assert extract_document('.docx',b.getvalue())[0]['quantity']==3
    pdf=PdfWriter();pdf.add_blank_page(100,100);b=io.BytesIO();pdf.write(b)
    assert extract_document('.pdf',b.getvalue())==[]

def test_upload_validation_review_and_no_mutation(system):
    from openpyxl import Workbook
    c,*_=system
    assert c.post('/api/upload',files={'file':('x.exe',b'xxx')}).status_code==422
    assert c.post('/api/upload',files={'file':('x.pdf',b'not-pdf')}).status_code==422
    assert c.post('/api/upload',files={'file':('x.pdf',b'x'*(5*1024*1024+70000))}).status_code==413
    wb=Workbook();wb.active.append(['DEMO-B16',2]);b=io.BytesIO();wb.save(b)
    response=c.post('/api/upload',files={'file':('a.xlsx',b.getvalue())})
    assert response.status_code==200,response.text
    assert response.json()['requires_review']
    assert c.get('/api/cart').json()['items']==[]

def test_photo_without_model_and_corrupt_image():
    from PIL import Image
    b=io.BytesIO();Image.new('RGB',(20,20)).save(b,format='JPEG')
    assert 'требует ключа' in extract('x.jpg',b.getvalue())['warning']
    with pytest.raises(Exception):extract('x.jpg',b'bad')

def test_model_tool_args_validated(system):
    _,_,_,app=system
    sid=app.state.cart.session()['id']
    with pytest.raises(ValueError):app.state.agent.tool(sid,'request_cart_add',{'product_id':'demo-2','quantity':True})
    with pytest.raises(ValueError):app.state.agent.tool(sid,'request_cart_add',{'product_id':'demo-2','quantity':1,'confirmed':True})
    with pytest.raises(ValueError):app.state.agent.tool(sid,'execute_code',{'code':'bad'})

def test_model_can_resolve_supplier_article(system):
    _,_,_,app=system
    sid=app.state.cart.session()['id']
    result=app.state.agent.tool(sid,'get_product',{'product_id':'027228'})
    assert result['products'][0]['id']=='515291'
    assert result['products'][0]['specs']['current'] is None

def test_responses_tool_loop(system):
    from agent.service import Agent
    _,cat,cart,_=system
    class Call:
        type='function_call'; name='get_product'; arguments='{"product_id":"demo-2"}';call_id='c1'
        def model_dump(self,**kw):return {'type':self.type,'name':self.name,'arguments':self.arguments,'call_id':self.call_id}
    responses=iter([SimpleNamespace(output=[Call()],output_text=''),SimpleNamespace(output=[],output_text='Данные получены.')])
    calls=[]
    def create(**kwargs):calls.append(kwargs);return next(responses)
    agent=Agent(cat,cart,SimpleNamespace(responses=SimpleNamespace(create=create)))
    r=agent.chat(cart.session()['id'],'DEMO-B16')
    assert r['products'][0]['id']=='demo-2' and r['mode']=='ai'
    assert calls[0]['store'] is False
    assert any(i.get('type')=='function_call_output' for i in calls[1]['input'])
